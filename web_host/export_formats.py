"""markdown → 纯文本 / docx 转换器（spec：2026-10-16-export-format-preference）。

markdown 是唯一真源：本模块只做表现层转换，不触碰 kernel。
to_text：标题去 #、行内标记还原纯文本，表格拍平为行；to_docx_bytes：#/##/###
映射 Heading 1/2/3，列表/表格/行内样式落地 Word；任何异常整体降级返回 None
（端点回落 txt）；filename_for：换扩展名。docx 全 run 设 w:eastAsia=宋体
（python-docx 只设 ascii 字体时中文会回落默认字体）。
"""
import io
import re

from docx import Document
from docx.oxml.ns import qn

_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`|\[[^\]\n]+\]\([^)\n]+\))")
_LINK = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_ULIST = re.compile(r"^\s*[-*]\s+(.*)$")
_OLIST = re.compile(r"^\s*\d+\.\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_TABLE_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def _strip_inline(text: str) -> str:
    out: list[str] = []
    for piece in _INLINE.split(text):
        if not piece:
            continue
        m = _LINK.match(piece)
        if m:
            out.append(f"{m.group(1)} ({m.group(2)})")
        elif piece.startswith("**") and piece.endswith("**"):
            out.append(piece[2:-2])
        elif piece.startswith("*") and piece.endswith("*"):
            out.append(piece[1:-1])
        elif piece.startswith("`") and piece.endswith("`"):
            out.append(piece[1:-1])
        else:
            out.append(piece)
    return "".join(out)


def to_text(markdown: str) -> str:
    lines_out: list[str] = []
    for line in markdown.splitlines():
        if _TABLE_SEP.match(line):
            continue
        if _TABLE_ROW.match(line):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            lines_out.append(" ".join(c for c in cells if c))
            continue
        m = _HEADING.match(line)
        if m:
            lines_out.append(_strip_inline(m.group(2).strip()))
            continue
        lines_out.append(_strip_inline(line))
    return "\n".join(lines_out)


def _songti(run) -> None:
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def _styled_run(paragraph, text: str, *, bold=False, italic=False, mono=False):
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if mono:
        run.font.name = "Consolas"
    _songti(run)
    return run


def _add_inline(paragraph, text: str) -> None:
    for piece in _INLINE.split(text):
        if not piece:
            continue
        m = _LINK.match(piece)
        if m:
            _styled_run(paragraph, f"{m.group(1)} ({m.group(2)})")
        elif piece.startswith("**") and piece.endswith("**"):
            _styled_run(paragraph, piece[2:-2], bold=True)
        elif piece.startswith("*") and piece.endswith("*"):
            _styled_run(paragraph, piece[1:-1], italic=True)
        elif piece.startswith("`") and piece.endswith("`"):
            _styled_run(paragraph, piece[1:-1], mono=True)
        else:
            _styled_run(paragraph, piece)


def to_docx_bytes(markdown: str) -> bytes | None:
    try:
        doc = Document()
        table_buffer: list[list[str]] = []

        def flush_table() -> None:
            if not table_buffer:
                return
            rows = len(table_buffer)
            cols = max(len(r) for r in table_buffer)
            table = doc.add_table(rows=rows, cols=cols)
            table.style = "Table Grid"
            for i, row in enumerate(table_buffer):
                for j in range(cols):
                    table.rows[i].cells[j].text = row[j] if j < len(row) else ""
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            _songti(r)
            table_buffer.clear()

        for line in markdown.splitlines():
            if _TABLE_ROW.match(line):
                if not _TABLE_SEP.match(line):
                    cells = [c.strip() for c in line.strip().strip("|").split("|")]
                    table_buffer.append(cells)
                continue
            flush_table()
            if not line.strip():
                continue
            m = _HEADING.match(line)
            if m:
                p = doc.add_heading("", level=min(len(m.group(1)), 4))
                _add_inline(p, _strip_inline(m.group(2).strip()))
                continue
            m = _ULIST.match(line)
            if m:
                _add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
                continue
            m = _OLIST.match(line)
            if m:
                _add_inline(doc.add_paragraph(style="List Number"), m.group(1))
                continue
            _add_inline(doc.add_paragraph(), line)
        flush_table()
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    except Exception:
        return None


def filename_for(base: str, fmt: str) -> str:
    stem = base.rsplit(".", 1)[0] if "." in base else base
    return f"{stem}.{fmt}"
