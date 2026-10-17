# tests/test_web_exports_api.py
import pytest
from fieldwork_agent.web_host.app import create_app
from httpx import ASGITransport, AsyncClient


@pytest.fixture
def app(tmp_path):
    from fieldwork_agent.kernel.repository.json_repo import JsonRepository
    from fieldwork_agent.web_host.deps import init_deps

    repo = JsonRepository(projects_dir=tmp_path / "projects")
    repo.create("p1")
    repo.save_proposal("p1", "# 标题\n\n## 一、背景\n\n内容",
                       kind="proposal", version="draft")
    _app = create_app(config=None, repo=repo)
    # ASGITransport 不触发 lifespan，手动初始化依赖
    init_deps(_app, config=None, repo=repo)
    return _app


@pytest.fixture
async def client(app):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as c:
        yield c


async def test_get_proposal(client):
    r = await client.get("/projects/p1/proposal")
    assert r.status_code == 200
    body = r.json()
    assert "# 标题" in body["markdown"]
    assert body["kind"] == "proposal" and body["version"] == "draft"


async def test_download_proposal(client):
    r = await client.get("/projects/p1/proposal/download")
    assert r.status_code == 200
    assert "text/markdown" in r.headers["content-type"]
    assert "attachment" in r.headers.get("content-disposition", "")
    assert "# 标题" in r.text


async def test_download_proposal_txt(client):
    r = await client.get("/projects/p1/proposal/download", params={"format": "txt"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/plain")
    cd = r.headers["content-disposition"]
    assert "attachment" in cd
    assert '.txt"' in cd
    assert "# 标题" not in r.text
    assert "一、背景" in r.text


async def test_download_proposal_docx(client):
    import io

    from docx import Document

    r = await client.get("/projects/p1/proposal/download", params={"format": "docx"})
    assert r.status_code == 200
    assert (
        r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    )
    cd = r.headers["content-disposition"]
    assert ".docx" in cd
    doc = Document(io.BytesIO(r.content))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert h1 == ["标题"]


async def test_download_proposal_invalid_format_422(client):
    r = await client.get("/projects/p1/proposal/download", params={"format": "pdf"})
    assert r.status_code == 422


async def test_download_proposal_docx_degrades_to_txt(client, monkeypatch):
    from fieldwork_agent.web_host.api import exports as exports_mod

    monkeypatch.setattr(exports_mod, "to_docx_bytes", lambda _md: None)
    r = await client.get("/projects/p1/proposal/download", params={"format": "docx"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/plain")
    assert '.txt"' in r.headers["content-disposition"]


async def test_download_proposal_explicit_md_matches_default(client):
    r_default = await client.get("/projects/p1/proposal/download")
    r_md = await client.get("/projects/p1/proposal/download", params={"format": "md"})
    assert r_md.status_code == r_default.status_code == 200
    assert r_md.headers["content-type"] == r_default.headers["content-type"]
    assert r_md.headers["content-disposition"] == r_default.headers["content-disposition"]
    assert r_md.content == r_default.content


async def test_list_proposals(client):
    r = await client.get("/projects/p1/proposals")
    assert r.status_code == 200
    assert len(r.json()) >= 1


async def test_proposal_404(client):
    r = await client.get("/projects/ghost/proposal")
    assert r.status_code == 404


async def test_get_proposal_by_id(client, app):
    repo = app.state.repo
    docs = repo.list_proposals("p1")
    r = await client.get(f"/projects/p1/proposals/{docs[0]['id']}")
    assert r.status_code == 200
    assert "# 标题" in r.json()["markdown"]
    assert r.json()["id"] == docs[0]["id"]


async def test_get_proposal_by_unknown_id_404(client):
    r = await client.get("/projects/p1/proposals/nope")
    assert r.status_code == 404


async def test_download_by_id(client, app):
    repo = app.state.repo
    docs = repo.list_proposals("p1")
    r = await client.get(f"/projects/p1/proposals/{docs[0]['id']}/download")
    assert r.status_code == 200
    assert "attachment" in r.headers.get("content-disposition", "")
    assert "# 标题" in r.text


async def test_download_by_id_txt(client, app):
    repo = app.state.repo
    docs = repo.list_proposals("p1")
    r = await client.get(f"/projects/p1/proposals/{docs[0]['id']}/download",
                         params={"format": "txt"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/plain")
    cd = r.headers["content-disposition"]
    assert "attachment" in cd
    assert '.txt"' in cd
    assert "# 标题" not in r.text
    assert "一、背景" in r.text


async def test_download_by_id_docx(client, app):
    import io

    from docx import Document

    repo = app.state.repo
    docs = repo.list_proposals("p1")
    r = await client.get(f"/projects/p1/proposals/{docs[0]['id']}/download",
                         params={"format": "docx"})
    assert r.status_code == 200
    assert (
        r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    )
    cd = r.headers["content-disposition"]
    assert ".docx" in cd
    doc = Document(io.BytesIO(r.content))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert h1 == ["标题"]


async def test_list_created_at_uses_saved_at(client):
    """口径修正：save_proposal 落盘的是 saved_at，列表不再恒 None。"""
    r = await client.get("/projects/p1/proposals")
    assert r.status_code == 200
    items = r.json()
    assert any(i.get("id", "").startswith("prop_") and i.get("created_at")
               for i in items), f"prop_ 记录的 created_at 应来自 saved_at: {items}"
