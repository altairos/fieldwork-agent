"""web_host/export_formats.py 单测：txt 拍平、docx 结构、降级 None、filename_for。

spec：docs/superpowers/specs/2026-10-16-export-format-preference-design.md
"""
import io

from docx import Document
from docx.oxml.ns import qn

from fieldwork_agent.web_host import export_formats as ef

MD = """# 研究方案标题

## 一、背景

调研 **粗体** 与 *斜体* 与 `代码` 与 [链接文本](https://example.com/a)。

- 要点甲
- 要点乙

1. 步骤一
2. 步骤二

| 项目 | 金额 |
| --- | --- |
| 住宿 | 1800 |
| 交通 | 900 |
"""


def test_to_text_strips_heading_and_inline():
    out = ef.to_text(MD)
    assert "研究方案标题" in out
    assert "#" not in out
    assert "**" not in out
    assert "*" not in out
    assert "`" not in out


def test_to_text_keeps_list_prefixes():
    out = ef.to_text(MD)
    assert "- 要点甲" in out
    assert "1. 步骤一" in out


def test_to_text_link_keeps_url():
    out = ef.to_text(MD)
    assert "链接文本 (https://example.com/a)" in out


def test_to_text_table_flattens():
    out = ef.to_text(MD)
    assert "|" not in out
    assert "---" not in out
    assert "项目 金额" in out
    assert "住宿 1800" in out


def test_to_docx_bytes_roundtrip():
    b = ef.to_docx_bytes(MD)
    assert b is not None
    assert b[:2] == b"PK"
    doc = Document(io.BytesIO(b))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert h1 == ["研究方案标题"]
    assert h2 == ["一、背景"]
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    numbers = [p.text for p in doc.paragraphs if p.style.name == "List Number"]
    assert bullets == ["要点甲", "要点乙"]
    assert numbers == ["步骤一", "步骤二"]
    assert doc.tables[0].style.name == "Table Grid"
    assert doc.tables[0].rows[0].cells[0].text == "项目"
    assert doc.tables[0].rows[2].cells[1].text == "900"
    bold_runs = [r for p in doc.paragraphs for r in p.runs if r.bold]
    assert bold_runs[0].text == "粗体"
    ea = bold_runs[0]._element.get_or_add_rPr().get_or_add_rFonts().get(qn("w:eastAsia"))
    assert ea == "宋体"


def test_to_docx_bytes_degrades_to_none(monkeypatch):
    def _boom(*args, **kwargs):
        raise RuntimeError("docx 渲染炸了")

    monkeypatch.setattr(ef, "Document", _boom)
    assert ef.to_docx_bytes(MD) is None


def test_filename_for():
    assert ef.filename_for("proposal_draft_v2.md", "md") == "proposal_draft_v2.md"
    assert ef.filename_for("proposal_draft_v2.md", "txt") == "proposal_draft_v2.txt"
    assert ef.filename_for("proposal_draft_v2.md", "docx") == "proposal_draft_v2.docx"
    assert ef.filename_for("no_ext", "txt") == "no_ext.txt"
